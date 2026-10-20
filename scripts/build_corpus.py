"""从 data/seeds 生成 data/raw 下的全部渲染产物。

用法：
    uv run python scripts/build_corpus.py --check      # 只跑自洽校验
    uv run python scripts/build_corpus.py              # 校验 + 全量渲染
    uv run python scripts/build_corpus.py --only xlsx  # 只渲染某类

设计说明：
- seeds 是事实源，raw 是可丢弃的渲染产物（可随时重建）。
- 数字自洽规则在渲染前**强制校验**，不通过直接退出，不生成半成品。
- DOCX 用 python-docx 生成，PDF 用 bundled LibreOffice 由 DOCX 转换，
  扫描版 PDF 用 pymupdf 把文本 PDF 栅格化后重新打包（无文本层）。
"""

from __future__ import annotations

import argparse
import copy
import json
import subprocess
import sys
from pathlib import Path
from typing import Any

import yaml

ROOT = Path(__file__).resolve().parents[1]
SEEDS = ROOT / "data" / "seeds"
RAW = ROOT / "data" / "raw"
BUILD = ROOT / "data" / "build"

# ---- 数字自洽规则的唯一来源（对应 docs/06_语料生成规格.md §6.1）----
ALLOWED_THICKNESS: dict[str, set[float]] = {
    "annealed": {4, 5, 6, 8, 10, 12},
    "toughened": {4, 5, 6, 8, 10, 12, 15, 19},
    "laminated": {6.38, 8.76, 10.76, 12.76},
    "igu": {4, 5, 6, 8, 10, 12},
}
IGU_CAVITY = {6, 9, 12, 16}
EXPERIENCE_PREFIXES = ("Usually", "In our experience", "Typically")
STANDARD_LIMITING_PHRASE = "subject to site measurement and NCC/DA requirements"


# ======================================================================
# 读取
# ======================================================================
def load_yaml(name: str) -> dict[str, Any]:
    with (SEEDS / name).open("r", encoding="utf-8") as fh:
        return yaml.safe_load(fh)


def load_csv(name: str) -> list[dict[str, str]]:
    import csv

    with (SEEDS / name).open("r", encoding="utf-8-sig", newline="") as fh:
        return list(csv.DictReader(fh))


def load_jsonl(name: str) -> list[dict[str, Any]]:
    rows = []
    with (SEEDS / name).open("r", encoding="utf-8") as fh:
        for line in fh:
            line = line.strip()
            if line:
                rows.append(json.loads(line))
    return rows


# ======================================================================
# 校验：数字自洽（§6）
# ======================================================================
class Checks:
    def __init__(self) -> None:
        self.errors: list[str] = []
        self.warnings: list[str] = []

    def fail(self, msg: str) -> None:
        self.errors.append(msg)

    def warn(self, msg: str) -> None:
        self.warnings.append(msg)


def check_capability(rows: list[dict[str, str]], ck: Checks) -> None:
    seen: set[tuple[str, float, str]] = set()
    for i, r in enumerate(rows, start=2):
        proc, shape = r["process"], r["shape"]
        th = float(r["thickness_mm"])
        key = (proc, th, shape)

        if proc not in ALLOWED_THICKNESS:
            ck.fail(f"capability 第 {i} 行：未知 process={proc}")
            continue
        if th not in ALLOWED_THICKNESS[proc]:
            ck.fail(
                f"capability 第 {i} 行：{proc} 不允许厚度 {th}mm"
                f"（允许 {sorted(ALLOWED_THICKNESS[proc])}）"
            )
        if key in seen:
            ck.fail(f"capability 第 {i} 行：唯一键重复 {key}")
        seen.add(key)

        if shape not in ("flat", "bent"):
            ck.fail(f"capability 第 {i} 行：未知 shape={shape}")
        if shape == "bent" and int(r["min_radius_mm"]) <= 0:
            ck.fail(f"capability 第 {i} 行：bent 但 min_radius_mm=0")
        if shape == "flat" and int(r["min_radius_mm"]) != 0:
            ck.fail(f"capability 第 {i} 行：flat 但 min_radius_mm≠0")

        # 面积与宽高自洽（允许 0.15 的四舍五入误差）
        w, h, area = int(r["max_width_mm"]), int(r["max_height_mm"]), float(r["max_area_m2"])
        computed = w * h / 1_000_000
        if abs(computed - area) > 0.15:
            ck.fail(
                f"capability 第 {i} 行：面积不自洽 {w}x{h}={computed:.2f}m2，"
                f"但 max_area_m2={area}"
            )
        if int(r["min_width_mm"]) > w or int(r["min_height_mm"]) > h:
            ck.fail(f"capability 第 {i} 行：最小值大于最大值")

        # 中空腔体
        if proc == "igu":
            if not r["igu_cavity_mm"]:
                ck.fail(f"capability 第 {i} 行：igu 缺 igu_cavity_mm")
            else:
                for cav in r["igu_cavity_mm"].split("/"):
                    if int(cav) not in IGU_CAVITY:
                        ck.fail(f"capability 第 {i} 行：腔体 {cav} 不在 {sorted(IGU_CAVITY)}")


def check_pricing(
    pricing: list[dict[str, str]],
    capability: list[dict[str, str]],
    leadtime: list[dict[str, str]],
    ck: Checks,
) -> None:
    cap_keys = {(r["process"], float(r["thickness_mm"]), r["shape"]) for r in capability}
    # 唯一键必须含 coating：同一厚度可有 none / low_e_soft 两种报价
    seen: set[tuple[str, float, str, str, str]] = set()

    for i, r in enumerate(pricing, start=2):
        proc = r["process"]
        th = float(r["thickness_mm"])
        shape = r["shape"]
        key = (proc, th, shape, r["variant"], r["coating"])

        if key in seen:
            ck.fail(f"pricing 第 {i} 行：唯一键重复 {key}")
        seen.add(key)

        if (proc, th, shape) not in cap_keys:
            ck.fail(
                f"pricing 第 {i} 行：能力矩阵中不存在 {proc}/{th}mm/{shape}"
            )

        # 成本价必须低于售价（用于验证脱敏后仍合理）
        unit, cost = float(r["unit_price_aud_m2"]), float(r["cost_price_aud_m2"])
        if not cost < unit:
            ck.fail(f"pricing 第 {i} 行：cost_price({cost}) 未低于 unit_price({unit})")

        if r["valid_from"] > r["valid_to"]:
            ck.fail(f"pricing 第 {i} 行：valid_from 晚于 valid_to")

        # 交期表必须覆盖该组合
        covered = any(
            lt["process"] == proc and float(lt["thickness_max_mm"]) >= th
            for lt in leadtime
        )
        if not covered:
            ck.fail(f"pricing 第 {i} 行：交期表未覆盖 {proc}/{th}mm")


def check_leadtime(rows: list[dict[str, str]], ck: Checks) -> None:
    for i, r in enumerate(rows, start=2):
        proc = r["process"]
        if proc not in ALLOWED_THICKNESS:
            ck.fail(f"leadtime 第 {i} 行：未知 process={proc}")
        if int(r["business_days"]) <= 0:
            ck.fail(f"leadtime 第 {i} 行：business_days 必须为正")
        if r["rush_available"].lower() == "false" and float(r["rush_surcharge_pct"]) != 0:
            ck.fail(f"leadtime 第 {i} 行：不可加急却有加急费")


def check_manuals(ck: Checks) -> None:
    for name in ("manual_tempering.yaml", "manual_iglazing.yaml"):
        doc = load_yaml(name)
        style = doc["doc"]["style"]
        blocks = _flatten_sections(doc["sections"])

        warnings = [b for b in blocks if b[0] == "warning"]
        if len(warnings) < 2:
            ck.fail(f"{name}：警告少于 2 条")
        levels = {b[2]["level"] for b in warnings}
        if not levels & {"DANGER", "WARNING"}:
            ck.fail(f"{name}：警告级别未包含 DANGER/WARNING")

        tables = [b for b in blocks if b[0] == "table"]
        for _, _, tbl, _depth in tables:
            if not tbl.get("caption"):
                ck.fail(f"{name}：表格缺 caption")
            if not tbl.get("columns") or not tbl.get("rows"):
                ck.fail(f"{name}：表格 {tbl.get('caption')} 缺 columns 或 rows")
            if any(len(row) != len(tbl["columns"]) for row in tbl["rows"]):
                ck.fail(f"{name}：表格 {tbl.get('caption')} 列数不一致")

        # 手册里出现的厚度必须合法
        if style == "checklist":
            for _, _, tbl, _depth in tables:
                cols = [c.lower() for c in tbl["columns"]]
                if cols and "thickness" in cols[0]:
                    for row in tbl["rows"]:
                        th = float(row[0])
                        if th not in ALLOWED_THICKNESS["toughened"]:
                            ck.fail(f"{name}：表格出现非法厚度 {th}mm")

        print(f"  {name:<24} style={style:<10} sections={len(doc['sections'])} "
              f"blocks={len(blocks)} tables={len(tables)} warnings={len(warnings)}")


def check_standards(ck: Checks) -> None:
    data = load_yaml("standards.yaml")
    if not data.get("disclaimer"):
        ck.fail("standards.yaml：缺 disclaimer")
    n = 0
    codes = []
    for std in data["standards"]:
        codes.append(std["code"])
        for sheet in std["sheets"]:
            n += 1
            ans = " ".join(sheet.get("answer_en", "").split()).lower()
            if STANDARD_LIMITING_PHRASE.lower() not in ans:
                ck.fail(f"standards {std['code']} / {sheet['topic']}：answer_en 缺限定语")
            if not sheet.get("section"):
                ck.fail(f"standards {std['code']} / {sheet['topic']}：缺 section")
            # 不允许出现具体条款号（形如 5.3.2）
            import re

            if re.search(r"\b\d+\.\d+\.\d+\b", sheet.get("section", "")):
                ck.fail(f"standards {std['code']}：section 出现具体条款号")
    print(f"  standards.yaml           codes={len(codes)} sheets={n}")


def starts_with_experience_prefix(text: str) -> bool:
    """经验类回答必须以对冲措辞开头（对应 constants.EXPERIENCE_PREFIX）。"""
    return text.strip().startswith(EXPERIENCE_PREFIXES)


def check_experience(ck: Checks) -> None:
    data = load_yaml("experience.yaml")
    entries = data["entries"]
    ids = [e["entry_id"] for e in entries]
    if len(ids) != len(set(ids)):
        ck.fail("experience.yaml：entry_id 重复")

    tag_count: dict[str, int] = {}
    for e in entries:
        ans = e["answer_en"].strip()
        if not starts_with_experience_prefix(ans):
            ck.fail(f"experience {e['entry_id']}：answer_en 未以 {EXPERIENCE_PREFIXES} 开头")
        if not e.get("confirmed_by") or not e.get("confirmed_at"):
            ck.fail(f"experience {e['entry_id']}：缺 confirmed_by / confirmed_at")
        if "$" in ans:
            ck.fail(f"experience {e['entry_id']}：出现价格符号（价格只能来自 pricing.csv）")
        if " days" in ans.lower():
            ck.warn(f"experience {e['entry_id']}：出现 ' days'，请确认没有写死交期")
        for t in e.get("tags", []):
            tag_count[t] = tag_count.get(t, 0) + 1

    expected_tags = {"spec", "config", "selection", "after_sales", "logistics"}
    missing = expected_tags - set(tag_count)
    if missing:
        ck.fail(f"experience.yaml：缺少意图覆盖 {missing}")
    print(f"  experience.yaml          entries={len(entries)} 分布={tag_count}")


def check_tickets(ck: Checks) -> None:
    rows = load_jsonl("tickets.jsonl")
    valid_intents = {"spec", "config", "selection", "after_sales", "logistics", "other"}
    ids = set()
    for r in rows:
        if r["intent"] not in valid_intents:
            ck.fail(f"tickets {r['ticket_id']}：未知 intent={r['intent']}")
        if r["ticket_id"] in ids:
            ck.fail(f"tickets：ticket_id 重复 {r['ticket_id']}")
        ids.add(r["ticket_id"])
        if r["date"] >= "2026-01-01":
            ck.fail(f"tickets {r['ticket_id']}：日期需早于系统上线 2026-01-01")
    esc = sum(1 for r in rows if r["escalated"])
    print(f"  tickets.jsonl            rows={len(rows)} escalated={esc}")


def check_manifest(ck: Checks) -> None:
    data = load_yaml("docs_manifest.yaml")
    docs = data["documents"]
    ids = [d["doc_id"] for d in docs]
    if len(ids) != len(set(ids)):
        ck.fail("docs_manifest.yaml：doc_id 重复")
    allowed_auth = {"standard", "manual", "experience"}
    for d in docs:
        if d["authority"] not in allowed_auth:
            ck.fail(f"manifest {d['doc_id']}：未知 authority={d['authority']}")
        if "version" in Path(d["file"]).name:
            ck.fail(f"manifest {d['doc_id']}：文件名不允许带版本号")
    print(f"  docs_manifest.yaml       documents={len(docs)}")


def _flatten_sections(sections: list[dict[str, Any]], depth: int = 1) -> list[tuple]:
    """把嵌套的 sections 摊平为有序块：('heading', number, title, depth) 等。"""
    out: list[tuple] = []
    for sec in sections:
        out.append(("heading", sec["number"], sec["title"], depth))
        for p in sec.get("paragraphs", []) or []:
            out.append(("paragraph", p, None, depth))
        for w in sec.get("warnings", []) or []:
            out.append(("warning", None, w, depth))
        for t in sec.get("tables", []) or []:
            out.append(("table", None, t, depth))
        if sec.get("subsections"):
            out.extend(_flatten_sections(sec["subsections"], depth + 1))
    return out


def run_checks() -> Checks:
    ck = Checks()
    print("数字自洽校验")
    capability = load_csv("capability.csv")
    pricing = load_csv("pricing.csv")
    leadtime = load_csv("leadtime.csv")

    check_capability(capability, ck)
    print(f"  capability.csv           rows={len(capability)}")
    check_leadtime(leadtime, ck)
    print(f"  leadtime.csv             rows={len(leadtime)}")
    check_pricing(pricing, capability, leadtime, ck)
    print(f"  pricing.csv              rows={len(pricing)}")
    check_manuals(ck)
    check_standards(ck)
    check_experience(ck)
    check_tickets(ck)
    check_manifest(ck)
    return ck


# ======================================================================
# 渲染：XLSX
# ======================================================================
def _write_xlsx(path: Path, headers: list[str], rows: list[list[Any]], title: str) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "data"

    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(vertical="center")

    for row in rows:
        ws.append(row)

    ws.freeze_panes = "A2"
    for idx, width in enumerate(_auto_width(headers, rows), start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width

    wb.properties.title = title
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def _auto_width(headers: list[str], rows: list[list[Any]]) -> list[int]:
    widths = []
    for i, h in enumerate(headers):
        longest = len(str(h))
        for r in rows:
            longest = max(longest, len(str(r[i])))
        widths.append(min(max(longest + 2, 10), 60))
    return widths


def render_tables() -> None:
    imports = load_yaml("manual_tempering.yaml")  # noqa: F841  (仅为保持加载顺序一致)

    # capability
    headers = ["Process", "Shape", "Thickness (mm)", "Max width (mm)", "Max height (mm)",
               "Max area (m2)", "Min width (mm)", "Min height (mm)", "Tolerance (mm)",
               "Min radius (mm)", "IGU cavity (mm)", "Notes"]
    rows = []
    for r in load_csv("capability.csv"):
        rows.append([r["process"], r["shape"], float(r["thickness_mm"]),
                     int(r["max_width_mm"]), int(r["max_height_mm"]), float(r["max_area_m2"]),
                     int(r["min_width_mm"]), int(r["min_height_mm"]), float(r["tolerance_mm"]),
                     int(r["min_radius_mm"]), r["igu_cavity_mm"] or None, r["notes"] or None])
    _write_xlsx(RAW / "tables" / "capability.xlsx", headers, rows,
                "Glass Capability Matrix 2026")
    print("  -> tables/capability.xlsx")

    # pricing
    headers = ["Process", "Shape", "Thickness (mm)", "Variant", "Coating",
               "Unit price (AUD/m2, incl GST)", "Min charge (AUD)", "MOQ (m2)",
               "Valid from", "Valid to", "Cost price (AUD/m2, internal)"]
    rows = []
    for r in load_csv("pricing.csv"):
        rows.append([r["process"], r["shape"], float(r["thickness_mm"]), r["variant"] or None,
                     r["coating"], float(r["unit_price_aud_m2"]), float(r["min_charge_aud"]),
                     float(r["moq_m2"]), r["valid_from"], r["valid_to"],
                     float(r["cost_price_aud_m2"])])
    _write_xlsx(RAW / "tables" / "pricing.xlsx", headers, rows, "Price List 2026")

    # leadtime
    headers = ["Process", "Thickness up to (mm)", "Quantity up to",
               "Business days", "Rush available", "Rush surcharge (%)", "Note"]
    rows = []
    for r in load_csv("leadtime.csv"):
        rows.append([r["process"], float(r["thickness_max_mm"]), int(r["qty_max"]),
                     int(r["business_days"]), r["rush_available"] == "true",
                     float(r["rush_surcharge_pct"]), r["note"] or None])
    _write_xlsx(RAW / "tables" / "leadtime.xlsx", headers, rows, "Standard Lead Times")
    print("  -> tables/pricing.xlsx, tables/leadtime.xlsx")


# ======================================================================
# 渲染：DOCX
# ======================================================================
def _build_docx(seed_name: str, out_path: Path) -> None:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    doc_data = load_yaml(seed_name)
    meta = doc_data["doc"]

    document = Document()

    # 标题页信息
    document.add_heading(meta["title"], level=0)
    info = document.add_paragraph()
    info.alignment = WD_ALIGN_PARAGRAPH.LEFT
    info.add_run(f"Revision {meta['revision']}    Language: {meta['language']}    "
                 f"Model: {meta['equipment']['model']}").italic = True
    document.add_paragraph()

    for block in _flatten_sections(doc_data["sections"]):
        kind = block[0]
        if kind == "heading":
            _, number, title, depth = block
            level = min(1 + (depth - 1), 3)
            document.add_heading(f"{number}  {title}", level=level)
        elif kind == "paragraph":
            document.add_paragraph(block[1])
        elif kind == "warning":
            w = block[2]
            p = document.add_paragraph()
            run = p.add_run(f"{w['level']}: ")
            run.bold = True
            p.add_run(w["text"])
        elif kind == "table":
            tbl = block[2]
            cap = document.add_paragraph()
            cap.add_run(tbl["caption"]).italic = True
            cap.runs[0].font.size = Pt(9)
            table = document.add_table(rows=1, cols=len(tbl["columns"]))
            table.style = "Table Grid"
            for i, col in enumerate(tbl["columns"]):
                cell = table.rows[0].cells[i]
                cell.text = ""
                cell.paragraphs[0].add_run(str(col)).bold = True
            for row in tbl["rows"]:
                cells = table.add_row().cells
                for i, val in enumerate(row):
                    cells[i].text = str(val)
            document.add_paragraph()

    out_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(out_path)


def render_docx() -> list[Path]:
    BUILD.mkdir(parents=True, exist_ok=True)
    made = []
    for seed, name in (("manual_tempering.yaml", "tempering"),
                       ("manual_iglazing.yaml", "iglazing")):
        # 中间产物放 data/build/，最终产物放 data/raw/
        build_path = BUILD / f"{name}.docx"
        _build_docx(seed, build_path)
        made.append(build_path)
        print(f"  -> build/{name}.docx")
    # tempering.docx 是 manifest 登记的真实文档，复制到 raw
    import shutil

    (RAW / "manuals").mkdir(parents=True, exist_ok=True)
    shutil.copy2(BUILD / "tempering.docx", RAW / "manuals" / "tempering.docx")
    print("  -> manuals/tempering.docx")
    return made


# ======================================================================
# 渲染：PDF（DOCX -> PDF，再栅格化成扫描版）
# ======================================================================
def _libreoffice_paths() -> tuple[str, str]:
    import os

    node = os.environ.get("LO_NODE", "")
    cli = os.environ.get("LO_CLI", "")
    if not node or not cli:
        raise SystemExit(
            "缺少 LibreOffice Kit 路径。请设置环境变量 LO_NODE / LO_CLI，或参见 README。"
        )
    return node, cli


def render_pdf() -> None:
    node, cli = _libreoffice_paths()
    (RAW / "manuals").mkdir(parents=True, exist_ok=True)

    for name in ("tempering", "iglazing"):
        src = BUILD / f"{name}.docx"
        out = RAW / "manuals" / f"{name}.pdf"
        # LibreOffice Kit 拒绝覆盖已存在文件，先删以保证幂等
        out.unlink(missing_ok=True)
        subprocess.run(
            [node, cli, "convert", "--input", str(src), "--output", str(out)],
            check=True,
            capture_output=True,
            text=True,
        )
        print(f"  -> manuals/{name}.pdf")


def render_scan_pdf(pages: int = 4, dpi: int = 120, jpeg_quality: int = 72) -> None:
    import pymupdf

    src = RAW / "manuals" / "tempering.pdf"
    out = RAW / "manuals" / "tempering_scan.pdf"

    doc = pymupdf.open(src)
    new = pymupdf.open()
    for i in range(min(pages, doc.page_count)):
        pix = doc[i].get_pixmap(dpi=dpi)
        jpeg = pix.tobytes("jpeg", jpg_quality=jpeg_quality)
        page = new.new_page(width=pix.width, height=pix.height)
        page.insert_image(pymupdf.Rect(0, 0, pix.width, pix.height), stream=jpeg)
    new.save(out, deflate=True)
    new.close()
    doc.close()

    # 验证：确实没有文本层
    check = pymupdf.open(out)
    text = "".join(check[i].get_text() for i in range(check.page_count))
    check.close()
    if text.strip():
        raise SystemExit(f"扫描版 PDF 仍含文本层（{len(text)} 字符），栅格化失败")
    size_mb = out.stat().st_size / 1024 / 1024
    print(f"  -> manuals/tempering_scan.pdf（{pages} 页，无文本层，{size_mb:.1f} MB，已验证）")


# ======================================================================
# 渲染：Markdown
# ======================================================================
def _slug(code: str) -> str:
    return (code.lower().replace("as/nzs", "as-nzs").replace(" ", "")
            .replace(".", "-").replace("/", "-"))


def render_standards_md() -> None:
    data = load_yaml("standards.yaml")
    outdir = RAW / "standards"
    outdir.mkdir(parents=True, exist_ok=True)

    for std in data["standards"]:
        slug = _slug(std["code"])
        lines = [
            f"# {std['code']} - {std['title']}",
            "",
            f"> {data['disclaimer']}",
            "",
        ]
        meta_bits = [f"Edition: {std.get('edition', 'current edition')}"]
        lines += ["  ".join(meta_bits), ""]

        for sheet in std["sheets"]:
            lines += [
                f"## {sheet['topic']}",
                "",
                f"**Standard**: {std['code']} - {std['title']}",
                f"**Section**: {sheet['section']}",
                "",
                "**Summary**:",
                sheet["summary"].strip(),
                "",
                "**Customer answer**:",
                sheet["answer_en"].strip(),
                "",
            ]
        (outdir / f"{slug}.md").write_text("\n".join(lines), encoding="utf-8")
        print(f"  -> standards/{slug}.md")


def render_fieldbook_md() -> None:
    data = load_yaml("experience.yaml")
    outdir = RAW / "experience"
    outdir.mkdir(parents=True, exist_ok=True)

    lines = [
        "# Customer Service Fieldbook",
        "",
        "> Internal notes compiled from production experience. "
        "Not a standard. Verify anything safety related before quoting.",
        "",
        f"Entries: {len(data['entries'])}",
        "",
    ]
    for e in data["entries"]:
        lines += [
            f"## {e['entry_id']} - {e['topic']}",
            "",
            f"**Question**: {e['question']}",
            "",
            e["answer_en"].strip(),
            "",
            f"*Confirmed by {e['confirmed_by']} on {e['confirmed_at']} "
            f"(ticket {e['source_ticket']})*",
            "",
        ]
    (outdir / "fieldbook.md").write_text("\n".join(lines), encoding="utf-8")
    print(f"  -> experience/fieldbook.md（{len(data['entries'])} 条）")


def render_tickets() -> None:
    outdir = RAW / "tickets"
    outdir.mkdir(parents=True, exist_ok=True)
    rows = load_jsonl("tickets.jsonl")
    with (outdir / "tickets.jsonl").open("w", encoding="utf-8", newline="\n") as fh:
        for r in rows:
            fh.write(json.dumps(r, ensure_ascii=False) + "\n")
    print(f"  -> tickets/tickets.jsonl（{len(rows)} 条）")


# ======================================================================
# main
# ======================================================================
def main() -> int:
    parser = argparse.ArgumentParser(description="GlassDesk 语料构建")
    parser.add_argument("--check", action="store_true", help="只跑自洽校验，不渲染")
    parser.add_argument("--only", choices=["xlsx", "docx", "pdf", "md", "jsonl"],
                        help="只渲染某类产物")
    args = parser.parse_args()

    ck = run_checks()

    for w in ck.warnings:
        print(f"  [warn] {w}")

    if ck.errors:
        print("\n校验失败：")
        for e in ck.errors:
            print(f"  [FAIL] {e}")
        return 1

    print("\n校验通过。")
    if args.check:
        return 0

    print("\n渲染产物")
    try:
        if args.only in (None, "xlsx"):
            render_tables()
        if args.only in (None, "docx"):
            render_docx()
        if args.only in (None, "pdf"):
            render_pdf()
            render_scan_pdf()
        if args.only in (None, "md"):
            render_standards_md()
            render_fieldbook_md()
        if args.only in (None, "jsonl"):
            render_tickets()
    except subprocess.CalledProcessError as exc:
        print(f"\nLibreOffice 转换失败：{exc.stderr or exc}")
        return 1

    print("\n完成。")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
