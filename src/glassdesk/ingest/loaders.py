"""五类文件 → 统一 Block 列表。

原则：
1. **loader 只还原结构，不做切分**。表格还原成"一行一块"，段落还原成一块，
   标题带层级。怎么切是 chunker 的事。
2. **表格行自描述**。每一行拼成 ``列名=值 | 列名=值``，并把表标题前置。
   这样一行脱离表格也能被检索和理解——教学项目把表格拍平，正是死在这里。
3. **扫描件只标记不处理**。文本层为空时置 ``needs_ocr=True``，不抛异常。
"""

from __future__ import annotations

import json
import re
from pathlib import Path

from glassdesk.ingest.models import Block, BlockKind, LoadResult
from glassdesk.logger import get_logger

log = get_logger("ingest.loaders")

#: 文本层字符数低于此值即判定为扫描件
_OCR_MIN_CHARS = 40

_WARNING_PREFIXES = ("DANGER:", "WARNING:", "CAUTION:", "NOTE:")
_TABLE_CAPTION = re.compile(r"^Table\s+[\d\-.]+\S*", re.IGNORECASE)


# ======================================================================
# DOCX
# ======================================================================
def load_docx(path: Path) -> LoadResult:
    """按文档顺序遍历段落与表格，保留 Heading 层级。"""
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = Document(str(path))
    blocks: list[Block] = []
    last_caption: str | None = None

    def iter_body(doc):
        body = doc.element.body
        for child in body.iterchildren():
            if child.tag.endswith("}p"):
                yield Paragraph(child, doc)
            elif child.tag.endswith("}tbl"):
                yield Table(child, doc)

    for item in iter_body(document):
        if isinstance(item, Paragraph):
            text = item.text.strip()
            if not text:
                continue
            style = (item.style.name or "") if item.style is not None else ""
            # python-docx 的 add_heading(level=0) 用的是 "Title" 样式，
            # 不是 "Heading 0"。漏掉它会把文档标题当成正文段落。
            if style == "Title":
                blocks.append(Block(text=text, kind=BlockKind.HEADING, level=0))
                last_caption = None
                continue
            if style.startswith("Heading"):
                level = 0
                for ch in reversed(style):
                    if ch.isdigit():
                        level = int(ch)
                        break
                blocks.append(Block(text=text, kind=BlockKind.HEADING, level=level))
                last_caption = None
                continue
            if text.upper().startswith(_WARNING_PREFIXES):
                blocks.append(Block(text=text, kind=BlockKind.WARNING))
                continue
            if _TABLE_CAPTION.match(text):
                last_caption = text
                blocks.append(Block(text=text, kind=BlockKind.PARAGRAPH))
                continue
            blocks.append(Block(text=text, kind=BlockKind.PARAGRAPH))
        else:  # Table
            rows = [[(c.text or "").strip() for c in row.cells] for row in item.rows]
            if not rows:
                continue
            header, data = rows[0], rows[1:]
            caption = last_caption or ""
            for row in data:
                pairs = [
                    f"{header[i]}={row[i]}" for i in range(min(len(header), len(row)))
                    if row[i] != ""
                ]
                if not pairs:
                    continue
                prefix = f"{caption} — " if caption else ""
                blocks.append(
                    Block(
                        text=prefix + " | ".join(pairs),
                        kind=BlockKind.TABLE_ROW,
                        clause=caption or None,
                    )
                )
            last_caption = None

    return LoadResult(blocks=blocks)


# ======================================================================
# PDF
# ======================================================================
def load_pdf(path: Path) -> LoadResult:
    """文本型 PDF：先抽表格，再把剩余文本按块抽出（避免表格文字重复计数）。"""
    import pymupdf

    doc = pymupdf.open(str(path))
    blocks: list[Block] = []
    total_chars = 0

    for page_index, page in enumerate(doc, start=1):
        table_rects = []
        try:
            found = page.find_tables()
            tables = list(found.tables)
        except Exception:  # 老版本或畸形页面，退化到纯文本
            tables = []

        for t_index, table in enumerate(tables, start=1):
            table_rects.append(pymupdf.Rect(table.bbox))
            try:
                rows = table.extract()
            except Exception:
                continue
            rows = [[(c or "").strip() for c in row] for row in rows if row]
            if len(rows) < 2:
                continue
            header, data = rows[0], rows[1:]
            caption = f"page {page_index} table {t_index}"
            for row in data:
                pairs = [
                    f"{header[i]}={row[i]}" for i in range(min(len(header), len(row)))
                    if row[i] != ""
                ]
                if not pairs:
                    continue
                blocks.append(
                    Block(
                        text=f"{caption} — " + " | ".join(pairs),
                        kind=BlockKind.TABLE_ROW,
                        page_no=page_index,
                        clause=caption,
                    )
                )

        for raw in page.get_text("blocks"):
            x0, y0, x1, y1, text = raw[0], raw[1], raw[2], raw[3], raw[4]
            rect = pymupdf.Rect(x0, y0, x1, y1)
            if any(rect.intersects(tr) for tr in table_rects):
                continue
            text = " ".join(text.split())
            if not text:
                continue
            total_chars += len(text)
            if text.upper().startswith(_WARNING_PREFIXES):
                kind = BlockKind.WARNING
            elif text.startswith(("1 ", "2 ", "3 ", "4 ", "5 ", "6 ")) and len(text) < 80:
                kind = BlockKind.HEADING
            else:
                kind = BlockKind.PARAGRAPH
            blocks.append(Block(text=text, kind=kind, page_no=page_index))

    page_count = doc.page_count
    doc.close()

    needs_ocr = page_count > 0 and total_chars < _OCR_MIN_CHARS
    if needs_ocr:
        log.warning("{} 文本层为空（{} 页），需要 OCR，M1a 阶段跳过", path.name, page_count)

    return LoadResult(blocks=blocks, needs_ocr=needs_ocr)


# ======================================================================
# XLSX
# ======================================================================
def load_xlsx(path: Path) -> LoadResult:
    """单 sheet、首行表头。每行一块，自描述。"""
    from openpyxl import load_workbook

    wb = load_workbook(str(path), read_only=True, data_only=True)
    ws = wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()

    if not rows:
        return LoadResult()

    header = [str(c).strip() if c is not None else "" for c in rows[0]]
    blocks: list[Block] = []
    for row in rows[1:]:
        pairs = [
            f"{header[i]}={row[i]}"
            for i in range(min(len(header), len(row)))
            if row[i] is not None and str(row[i]).strip() != ""
        ]
        if not pairs:
            continue
        blocks.append(Block(text=" | ".join(pairs), kind=BlockKind.TABLE_ROW,
                            clause=ws.title))

    return LoadResult(blocks=blocks)


# ======================================================================
# Markdown
# ======================================================================
def load_md(path: Path) -> LoadResult:
    """按 ``##`` 分节。每节内容合成一块，交给 chunker 整节处理。"""
    text = path.read_text(encoding="utf-8")
    blocks: list[Block] = []
    current_heading: str | None = None
    buffer: list[str] = []

    def flush() -> None:
        if not buffer:
            return
        body = " ".join(" ".join(buffer).split())
        if body:
            blocks.append(
                Block(text=body, kind=BlockKind.PARAGRAPH, clause=current_heading)
            )
        buffer.clear()

    for line in text.splitlines():
        stripped = line.strip()
        if stripped.startswith("## "):
            flush()
            current_heading = stripped[3:].strip()
            blocks.append(Block(text=current_heading, kind=BlockKind.HEADING, level=2))
            continue
        if stripped.startswith("# "):
            flush()
            blocks.append(Block(text=stripped[2:].strip(), kind=BlockKind.HEADING, level=1))
            continue
        if stripped.startswith(">"):  # 引用块（免责声明）不入库
            continue
        buffer.append(stripped)

    flush()
    return LoadResult(blocks=blocks)


# ======================================================================
# JSONL
# ======================================================================
def load_jsonl(path: Path) -> LoadResult:
    """一行一条记录，拼成自描述的键值串。"""
    blocks: list[Block] = []
    with path.open("r", encoding="utf-8") as fh:
        for line in fh:
            line = line.strip()
            if not line:
                continue
            record = json.loads(line)
            pairs = [f"{k}={v}" for k, v in record.items()
                     if v is not None and str(v).strip() != ""]
            if not pairs:
                continue
            blocks.append(Block(text=" | ".join(pairs), kind=BlockKind.RECORD))
    return LoadResult(blocks=blocks)


# ======================================================================
# 分发
# ======================================================================
LOADERS = {
    ".docx": load_docx,
    ".pdf": load_pdf,
    ".xlsx": load_xlsx,
    ".md": load_md,
    ".jsonl": load_jsonl,
}

SUPPORTED_SUFFIXES: tuple[str, ...] = tuple(LOADERS)


def load_file(path: str | Path) -> LoadResult:
    """按扩展名分发。不支持的类型抛 ValueError，不静默跳过。"""
    p = Path(path)
    suffix = p.suffix.lower()
    loader = LOADERS.get(suffix)
    if loader is None:
        raise ValueError(
            f"不支持的文件类型 {suffix}（{p.name}）；已支持：{', '.join(SUPPORTED_SUFFIXES)}"
        )
    if not p.is_file():
        raise FileNotFoundError(f"找不到文件：{p}")
    log.info("解析 {}（{}）", p.name, suffix)
    return loader(p)
