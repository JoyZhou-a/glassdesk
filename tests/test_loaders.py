"""Loader 的单元测试：全部在 tmp_path 里现造文件，不依赖 data/raw。"""

from __future__ import annotations

import json
from pathlib import Path

import pytest

from glassdesk.ingest.loaders import load_file
from glassdesk.ingest.models import BlockKind, Chunk, sha256_text


# ======================================================================
# 造文件的小工具
# ======================================================================
@pytest.fixture
def docx_file(tmp_path: Path) -> Path:
    from docx import Document

    doc = Document()
    doc.add_heading("TF-2448 Manual", level=0)
    doc.add_heading("1  Safety", level=1)
    doc.add_paragraph("Only trained operators may run a cycle.")
    p = doc.add_paragraph()
    p.add_run("DANGER: ").bold = True
    p.add_run("Never open the door while pressurised.")
    doc.add_paragraph("Table 4-1 Recommended heating time")
    table = doc.add_table(rows=1, cols=3)
    table.rows[0].cells[0].text = "Thickness (mm)"
    table.rows[0].cells[1].text = "Heating time (s/mm)"
    table.rows[0].cells[2].text = "Quench (kPa)"
    for values in (("4", "38", "12"), ("6", "42", "18")):
        cells = table.add_row().cells
        for i, v in enumerate(values):
            cells[i].text = v
    out = tmp_path / "manual.docx"
    doc.save(out)
    return out


@pytest.fixture
def xlsx_file(tmp_path: Path) -> Path:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "data"
    ws.append(["Process", "Thickness (mm)", "Max width (mm)"])
    ws.append(["toughened", 6, 2440])
    ws.append(["toughened", 8, 2440])
    out = tmp_path / "cap.xlsx"
    wb.save(out)
    return out


@pytest.fixture
def md_file(tmp_path: Path) -> Path:
    out = tmp_path / "std.md"
    out.write_text(
        "# AS 1288 - Glass in buildings\n"
        "\n"
        "> Synthetic summary for demonstration. Not the standard text.\n"
        "\n"
        "## Safety glazing in doors\n"
        "\n"
        "**Standard**: AS 1288\n"
        "\n"
        "Doors must be safety glazed.\n"
        "\n"
        "## Overhead glazing\n"
        "\n"
        "Overhead must be laminated.\n",
        encoding="utf-8",
    )
    return out


@pytest.fixture
def jsonl_file(tmp_path: Path) -> Path:
    out = tmp_path / "tickets.jsonl"
    rows = [
        {"ticket_id": "T-1", "intent": "spec", "escalated": False},
        {"ticket_id": "T-2", "intent": "after_sales", "escalated": True},
    ]
    out.write_text("\n".join(json.dumps(r) for r in rows) + "\n", encoding="utf-8")
    return out


@pytest.fixture
def text_pdf(tmp_path: Path) -> Path:
    import pymupdf

    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 100), "Section 2 Overview and Specifications", fontsize=12)
    page.insert_text((72, 130), "The furnace accepts flat glass from 4 mm to 19 mm.", fontsize=10)
    out = tmp_path / "text.pdf"
    doc.save(out)
    doc.close()
    return out


@pytest.fixture
def scan_pdf(tmp_path: Path) -> Path:
    """只有图片、没有文本层——模拟扫描件。"""
    import pymupdf

    doc = pymupdf.open()
    page = doc.new_page()
    pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 200, 200))
    pix.set_rect(pix.irect, (230, 230, 230))
    page.insert_image(pymupdf.Rect(72, 72, 400, 400), pixmap=pix)
    out = tmp_path / "scan.pdf"
    doc.save(out)
    doc.close()
    return out


# ======================================================================
# DOCX
# ======================================================================
def test_docx_preserves_heading_levels(docx_file: Path):
    result = load_file(docx_file)
    headings = [b for b in result.blocks if b.kind is BlockKind.HEADING]
    assert [h.level for h in headings] == [0, 1]
    assert headings[1].text == "1  Safety"


def test_docx_detects_warnings(docx_file: Path):
    result = load_file(docx_file)
    warnings = [b for b in result.blocks if b.kind is BlockKind.WARNING]
    assert len(warnings) == 1
    assert warnings[0].text.startswith("DANGER:")


def test_docx_table_row_is_self_describing(docx_file: Path):
    """核心断言：一行脱离表格也能读懂，且带表标题。"""
    result = load_file(docx_file)
    rows = [b for b in result.blocks if b.kind is BlockKind.TABLE_ROW]
    assert len(rows) == 2
    assert rows[0].text.startswith("Table 4-1 Recommended heating time — ")
    assert "Thickness (mm)=4" in rows[0].text
    assert "Heating time (s/mm)=38" in rows[0].text
    assert rows[0].clause == "Table 4-1 Recommended heating time"


def test_docx_has_no_empty_blocks(docx_file: Path):
    result = load_file(docx_file)
    assert all(b.text.strip() for b in result.blocks)


# ======================================================================
# XLSX
# ======================================================================
def test_xlsx_one_block_per_data_row(xlsx_file: Path):
    result = load_file(xlsx_file)
    rows = [b for b in result.blocks if b.kind is BlockKind.TABLE_ROW]
    assert len(rows) == 2
    assert "Process=toughened" in rows[0].text
    assert "Thickness (mm)=6" in rows[0].text


def test_xlsx_skips_header_only_sheet(tmp_path: Path):
    from openpyxl import Workbook

    wb = Workbook()
    wb.active.append(["A", "B"])
    out = tmp_path / "empty.xlsx"
    wb.save(out)
    assert len(load_file(out).blocks) == 0


# ======================================================================
# Markdown
# ======================================================================
def test_md_splits_by_h2_and_ignores_blockquote(md_file: Path):
    result = load_file(md_file)
    headings = [b for b in result.blocks if b.kind is BlockKind.HEADING]
    assert [h.text for h in headings] == [
        "AS 1288 - Glass in buildings",
        "Safety glazing in doors",
        "Overhead glazing",
    ]
    bodies = [b for b in result.blocks if b.kind is BlockKind.PARAGRAPH]
    assert bodies[0].clause == "Safety glazing in doors"
    # 引用块（免责声明）不应进入正文
    assert all("Synthetic summary" not in b.text for b in result.blocks)


# ======================================================================
# JSONL
# ======================================================================
def test_jsonl_one_record_per_line(jsonl_file: Path):
    result = load_file(jsonl_file)
    assert len(result.blocks) == 2
    assert result.blocks[1].kind is BlockKind.RECORD
    assert "ticket_id=T-2" in result.blocks[1].text


# ======================================================================
# PDF
# ======================================================================
def test_text_pdf_extracts_blocks(text_pdf: Path):
    result = load_file(text_pdf)
    assert result.needs_ocr is False
    assert result.blocks
    assert all(b.page_no == 1 for b in result.blocks)


def test_scanned_pdf_flags_needs_ocr(scan_pdf: Path):
    """扫描件不抛异常，只标记；M1a 阶段不处理 OCR。"""
    result = load_file(scan_pdf)
    assert result.needs_ocr is True
    assert result.blocks == []


# ======================================================================
# 分发与错误
# ======================================================================
def test_unsupported_suffix_raises(tmp_path: Path):
    p = tmp_path / "a.txt"
    p.write_text("x", encoding="utf-8")
    with pytest.raises(ValueError, match="不支持的文件类型"):
        load_file(p)


def test_missing_file_raises(tmp_path: Path):
    with pytest.raises(FileNotFoundError):
        load_file(tmp_path / "nope.docx")


# ======================================================================
# Chunk 模型
# ======================================================================
def test_chunk_rejects_empty_text():
    c = Chunk(chunk_id="x", doc_id="d", version=1, kind="child",
              block_kind="paragraph", ordinal=0, text="   |  |  ")
    assert c.is_valid() is False


def test_chunk_content_hash_is_stable():
    a = Chunk(chunk_id="a", doc_id="d", version=1, kind="child",
              block_kind="paragraph", ordinal=0, text="hello")
    b = Chunk(chunk_id="b", doc_id="d", version=1, kind="child",
              block_kind="paragraph", ordinal=1, text="hello")
    assert a.content_hash == b.content_hash == sha256_text("hello")
